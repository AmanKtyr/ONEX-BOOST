from django.shortcuts import render, redirect, get_object_or_404
from .forms import MemberForm, MedicalHistoryForm, EmergencyContactForm, MembershipHistoryForm, PersonalTrainerForm, AssignDietPlanForm, AssignWorkoutPlanForm
from .models import Member, MedicalHistory, EmergencyContact, MembershipHistory, PersonalTrainer, MembershipFreeze, AssignDietPlan, AssignWorkoutPlan, MemberWorkoutLog, MemberBodyMetric, MemberFitnessGoal
from apps.management.models import MembershipPlan
from apps.trainers.models import Trainer
from apps.billing.models import Payment
from django.forms import modelformset_factory
from django.contrib import messages
from django.core.paginator import Paginator
from django.db.models import Q, Sum, F, Prefetch
from django.http import JsonResponse
from django.core.serializers import serialize
from django.views.decorators.http import require_POST, require_GET
from django.db import IntegrityError, transaction
from apps.whatsapp.services import WhatsAppService
from apps.superadmin.notifications import (
    notify_membership_activated,
    notify_pt_assigned,
    notify_payment_submitted
)
from apps.superadmin.subscription_utils import check_resource_quota, get_gym_quota_status

from django.contrib.auth.decorators import login_required
from apps.login.decorators import custom_permission_required
from django.views.decorators.cache import never_cache
from django.utils import timezone
from datetime import timedelta, date
from decimal import Decimal, ROUND_HALF_UP
import logging
logger = logging.getLogger(__name__)


@require_POST
def check_mobile_number_registration(request):
    mobile_number = request.POST.get('mobile_number')
    gym = getattr(request, 'gym', None)
    is_registered = False
    registered_name = ""

    if mobile_number:
        member = Member.objects.filter(gym=gym, mobile_number=mobile_number).first()
        if member:
            is_registered = True
            registered_name = member.first_name # Assuming first_name is sufficient for display

    return JsonResponse({'is_registered': is_registered, 'registered_name': registered_name})


@never_cache
@login_required(login_url='login')
@custom_permission_required('add_member')
def add_new_member(request):
    gym = getattr(request, 'gym', None)
    can_add, quota_error, quota_status = check_resource_quota(gym, 'member')

    MedicalHistoryFormSet = modelformset_factory(MedicalHistory, form=MedicalHistoryForm, extra=1, can_delete=True)
    if request.method == 'POST':
        if not can_add:
            messages.error(request, quota_error)
            return redirect('member_list')

        member_form = MemberForm(request.POST, request.FILES)
        medical_formset = MedicalHistoryFormSet(request.POST, request.FILES, prefix='medical')
        emergency_form = EmergencyContactForm(request.POST, prefix='emergency')

        if member_form.is_valid() and medical_formset.is_valid() and emergency_form.is_valid():
            try:
                member = member_form.save(commit=False)
                member.gym = gym
                member.save()
                
                instances = medical_formset.save(commit=False)
                for instance in instances:
                    instance.member = member
                    instance.gym = gym
                    instance.save()
                
                medical_formset.save_m2m() 
                
                for form in medical_formset.deleted_forms:
                    if form.instance.pk:
                        form.instance.delete()

                emergency_contact = emergency_form.save(commit=False)
                emergency_contact.member = member
                emergency_contact.gym = gym
                emergency_contact.save()

                # Provision Member User Account
                user, raw_password = member.create_or_update_user_account()
                request.session['registration_credentials'] = {
                    'account_type': 'Member',
                    'title': 'Member Registered Successfully!',
                    'name': f"{member.first_name} {member.last_name}".strip(),
                    'id': member.member_id,
                    'username': member.member_id,
                    'mobile': member.mobile_number,
                    'email': member.email or '',
                    'password': raw_password,
                    'gym_name': gym.name if gym else 'ONEX BOOST Gym',
                    'login_url': request.build_absolute_uri('/login/'),
                }

                messages.success(request, 'Member added successfully!')
                return redirect('assign_membership_plan', member_id=member.id)
            except IntegrityError as e:
                if 'email' in str(e):
                    member_form.add_error('email', 'A member with this email already exists.')
                elif 'mobile_number' in str(e):
                    member_form.add_error('mobile_number', 'A member with this mobile number and relationship already exists.')
                else:
                    messages.error(request, 'An unexpected error occurred. Please try again.')

        else:
            logger.warning("Member form invalid. Errors: %s | Medical: %s | Emergency: %s",
                           member_form.errors, medical_formset.errors, emergency_form.errors)
    else:
        member_form = MemberForm()
        medical_formset = MedicalHistoryFormSet(queryset=MedicalHistory.objects.none(), prefix='medical')
        emergency_form = EmergencyContactForm(prefix='emergency')
    return render(request, 'members/add_new_member.html', {
        'form': member_form,
        'medical_formset': medical_formset,
        'emergency_form': emergency_form,
        'quota_status': quota_status,
        'quota_limit_reached': not can_add,
        'quota_error': quota_error,
    })

@never_cache
@login_required(login_url='login')
def member_profile(request, member_id):
    gym = getattr(request, 'gym', None)
    member = Member.objects.filter(id=member_id, gym=gym, is_deleted=False).first()
    if not member:
        messages.error(request, 'Member not found or has been removed.')
        return redirect('member_list')
    
    # Fetch both active and frozen memberships (excluding trashed)
    membership_histories = MembershipHistory.objects.filter(
        member=member, gym=gym, is_deleted=False
    ).order_by('-id')
    
    pt_member = PersonalTrainer.objects.select_related('trainer').filter(
        member=member, status='active', gym=gym, is_deleted=False
    ).order_by('-id')
    
    # The latest membership can be active or frozen
    latest_membership = membership_histories.first()
    
    # Determine if the plan is active
    is_plan_active = member.current_status == 'Active'
    
    payments = Payment.objects.filter(
        member=member,
        gym=gym,
        is_deleted=False
    ).exclude(
        membership_history__is_deleted=True
    ).exclude(
        personal_trainer__is_deleted=True
    ).order_by('-payment_date')

    # Calculate the total due amount for active memberships only
    membership_due_amount = membership_histories.filter(status='active').aggregate(
        total_due=Sum(F('total_amount') - F('paid_amount'))
    )['total_due'] or 0

    # Calculate the total due amount for personal training
    pt_due_amount = pt_member.filter(status='active').aggregate(
        total_due=Sum(F('total_amount') - F('paid_amount'))
    )['total_due'] or 0

    total_due_amount = membership_due_amount + pt_due_amount

    assigned_diet_plans = AssignDietPlan.objects.filter(member=member, gym=gym).order_by('-assigned_at')
    assigned_workout_plans = AssignWorkoutPlan.objects.filter(member=member, gym=gym).order_by('-assigned_at')
    workout_logs = MemberWorkoutLog.objects.filter(member=member, gym=gym).prefetch_related('exercises').order_by('-date', '-created_at')[:30]
    body_metrics = MemberBodyMetric.objects.filter(member=member, gym=gym).order_by('-date', '-created_at')[:30]
    fitness_goal = MemberFitnessGoal.objects.filter(member=member, gym=gym).first()

    return render(request, 'members/member_profile.html', {
        'member': member, 
        'membership_histories': membership_histories,
        'pt_member': pt_member,
        'membership_history': latest_membership,
        'due_amount': total_due_amount,
        'pt_invoices': pt_member,
        'payments': payments,
        'membership_plan': latest_membership.plan if latest_membership else None,
        'assigned_diet_plans': assigned_diet_plans,
        'assigned_workout_plans': assigned_workout_plans,
        'workout_logs': workout_logs,
        'body_metrics': body_metrics,
        'fitness_goal': fitness_goal,
        'is_plan_active': is_plan_active,
    })


@never_cache
@login_required(login_url='login')
@custom_permission_required('change_member')
def edit_member(request, member_id):
    gym = getattr(request, 'gym', None)
    member = Member.objects.filter(id=member_id, gym=gym).first()
    if not member:
        messages.error(request, 'Member not found.')
        return redirect('member_list')
    try:
        emergency_contact = member.emergency_contact
    except EmergencyContact.DoesNotExist:
        emergency_contact = None

    MedicalHistoryFormSet = modelformset_factory(MedicalHistory, form=MedicalHistoryForm, extra=1, can_delete=True)

    if request.method == 'POST':
        form = MemberForm(request.POST, request.FILES, instance=member)
        medical_formset = MedicalHistoryFormSet(request.POST, request.FILES, queryset=MedicalHistory.objects.filter(member=member, gym=gym), prefix='medical')
        emergency_form = EmergencyContactForm(request.POST, instance=emergency_contact, prefix='emergency')

        if form.is_valid() and medical_formset.is_valid() and emergency_form.is_valid():
            form.save()
            
            instances = medical_formset.save(commit=False)
            for instance in instances:
                instance.member = member
                instance.gym = gym
                instance.save()
            
            medical_formset.save_m2m()

            for form_in_formset in medical_formset.deleted_forms:
                if form_in_formset.instance.pk:
                    form_in_formset.instance.delete()

            emergency_contact_instance = emergency_form.save(commit=False)
            emergency_contact_instance.member = member
            emergency_contact_instance.gym = gym
            emergency_contact_instance.save()
            
            messages.success(request, 'Member details updated successfully!')
            return redirect('member_list')
        else:
            # For debugging purposes
            logger.warning("Edit member form invalid. Errors: %s | Medical: %s | Emergency: %s",
                           form.errors, medical_formset.errors, emergency_form.errors)

    else:
        form = MemberForm(instance=member)
        medical_formset = MedicalHistoryFormSet(queryset=MedicalHistory.objects.filter(member=member, gym=gym), prefix='medical')
        emergency_form = EmergencyContactForm(instance=emergency_contact, prefix='emergency')

    return render(request, 'members/edit_member.html', {
        'form': form,
        'medical_formset': medical_formset,
        'emergency_form': emergency_form,
        'member': member
    })


@never_cache
@login_required(login_url='login') 
@custom_permission_required('view_member')
def member_list(request):
    gym = getattr(request, 'gym', None)
    query = request.GET.get('q')
    status_filter = request.GET.get('status_filter')
    date_from = request.GET.get('date_from')
    date_to = request.GET.get('date_to')

    # Initial fetch with prefetching
    member_qs = Member.objects.filter(gym=gym, is_deleted=False).prefetch_related(
        Prefetch('membership_history',
                 queryset=MembershipHistory.objects.filter(is_deleted=False).select_related('plan').prefetch_related('freezes')),
        Prefetch('personal_trainer', queryset=PersonalTrainer.objects.filter(is_deleted=False).select_related('trainer')),
    )

    if query:
        query_parts = query.split()
        q_objects = Q()
        if len(query_parts) > 1:
            q_objects.add(Q(first_name__icontains=query_parts[0]) & Q(last_name__icontains=" ".join(query_parts[1:])), q_objects.OR)
        q_objects.add(Q(first_name__icontains=query) |
                      Q(last_name__icontains=query) |
                      Q(email__icontains=query) |
                      Q(mobile_number__icontains=query) |
                      Q(member_id__icontains=query), q_objects.OR)
        member_qs = member_qs.filter(q_objects).distinct()

    # Parse date filters
    date_from_obj = None
    date_to_obj = None
    if date_from:
        try:
            date_from_obj = date.fromisoformat(date_from)
        except ValueError:
            pass
    if date_to:
        try:
            date_to_obj = date.fromisoformat(date_to)
        except ValueError:
            pass

    # Process all members to compute status and expiry
    today = timezone.localdate()
    seven_days_from_now = today + timedelta(days=7)
    
    all_members_data = []
    for member in member_qs:
        histories = list(member.membership_history.all())
        latest = max(histories, key=lambda mh: mh.membership_start_date, default=None)
        first = min(histories, key=lambda mh: mh.membership_start_date, default=None)
        
        member.latest_membership_history = latest
        member.all_histories_list = histories
        member.join_date = first.membership_start_date if first else None
        
        end_date = latest.get_end_date() if latest else None
        member.computed_end_date = end_date
        
        # Determine Status Display
        if not latest:
            member.status_display = 'No Membership'
            member.computed_status = 'no_membership'
        elif end_date and end_date < today:
            member.status_display = 'Expired'
            member.computed_status = 'expired'
            if latest.status == 'active':
                latest.status = 'inactive'
                latest.save(update_fields=['status'])
        elif latest.status == 'freezed':
            member.status_display = 'Frozen'
            member.computed_status = 'freezed'
        elif end_date and today <= end_date <= seven_days_from_now:
            member.status_display = 'Expiring Soon'
            member.computed_status = 'expiring_soon'
        elif latest.membership_start_date > today:
            member.status_display = 'Upcoming'
            member.computed_status = 'upcoming'
        else:
            member.status_display = 'Active'
            member.computed_status = 'active'
            
        # Apply Smart Date Filtering
        if date_from_obj or date_to_obj:
            # Determine which date we are filtering by
            if status_filter in ['expired', 'expiring_soon']:
                relevant_date = end_date
            elif status_filter == 'upcoming':
                relevant_date = latest.membership_start_date if latest else None
            else:
                relevant_date = member.join_date
            
            if not relevant_date:
                continue
            if date_from_obj and relevant_date < date_from_obj:
                continue
            if date_to_obj and relevant_date > date_to_obj:
                continue

        all_members_data.append(member)

    # Apply Status Filter from computed data
    if status_filter:
        if status_filter == 'active':
            all_members_data = [m for m in all_members_data if m.computed_status in ['active', 'expiring_soon']]
        elif status_filter == 'expired':
            all_members_data = [m for m in all_members_data if m.computed_status == 'expired']
        elif status_filter == 'expiring_soon':
            all_members_data = [m for m in all_members_data if m.computed_status == 'expiring_soon']
        elif status_filter == 'freezed':
            all_members_data = [m for m in all_members_data if m.computed_status == 'freezed']
        elif status_filter == 'upcoming':
            all_members_data = [m for m in all_members_data if m.computed_status == 'upcoming']
        elif status_filter == 'no_membership':
            all_members_data = [m for m in all_members_data if m.computed_status == 'no_membership']
        elif status_filter == 'due_members':
            due_members = []
            for m in all_members_data:
                m_due = sum((mh.total_amount - mh.paid_amount) for mh in m.all_histories_list if mh.status == 'active')
                p_due = sum((pt.total_amount - pt.paid_amount) for pt in m.personal_trainer.all() if pt.status == 'active')
                if (m_due + p_due) > 0:
                    due_members.append(m)
            all_members_data = due_members

    # Apply Sorting
    if status_filter == 'expired':
        all_members_data.sort(key=lambda x: (x.computed_end_date or date.min), reverse=True)
    elif status_filter in ['expiring_soon', 'active']:
        all_members_data.sort(key=lambda x: (x.computed_end_date or date.max))
    elif status_filter == 'upcoming':
        all_members_data.sort(key=lambda x: x.latest_membership_history.membership_start_date)
    else:
        all_members_data.sort(key=lambda x: x.id, reverse=True)

    if request.GET.get('export') == 'excel':
        import openpyxl
        from openpyxl.styles import Font, Alignment
        from django.http import HttpResponse

        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="members.xlsx"'

        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = 'Members'

        columns = [
            'Photo', 'Full Name', 'Mobile Number', 'Joining Date',
            'Current/Last Plan', 'Due Amount', 'PT Plan Details', 'Full Address'
        ]

        for col_num, column_title in enumerate(columns, 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.value = column_title
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            column_letter = openpyxl.utils.get_column_letter(col_num)
            worksheet.column_dimensions[column_letter].width = 30

        for index, member in enumerate(all_members_data, 1):
            photo_url = request.build_absolute_uri(member.profile_picture.url) if member.profile_picture else 'No Photo'
            
            full_name = f"{member.first_name} {member.last_name}"
            
            joining_date = member.join_date.strftime('%d-%m-%Y') if member.join_date else ''

            latest = member.latest_membership_history
            plan_info = 'No Plan'
            if latest:
                plan_info = f"{latest.plan.title} ({latest.plan.get_duration_display()}) - From: {latest.membership_start_date.strftime('%d-%m-%Y')} To: {latest.get_end_date().strftime('%d-%m-%Y') if latest.get_end_date() else 'N/A'}"

            active_pt = member.personal_trainer.filter(status='active').first()
            pt_info = 'No PT'
            if active_pt:
                pt_info = f"Trainer: {active_pt.trainer.name} - From: {active_pt.pt_start_date.strftime('%d-%m-%Y')} To: {active_pt.get_end_date().strftime('%d-%m-%Y') if active_pt.get_end_date() else 'N/A'}"

            address_parts = [member.address, member.city, member.state, member.pincode]
            full_address = ', '.join(part for part in address_parts if part)

            histories = member.all_histories_list
            membership_due_amount = sum((mh.total_amount - mh.paid_amount) for mh in histories if mh.status == 'active')
            pt_due_amount = sum((pt.total_amount - pt.paid_amount) for pt in member.personal_trainer.all() if pt.status == 'active')
            total_due = membership_due_amount + pt_due_amount

            row_data = [
                photo_url,
                full_name,
                member.mobile_number,
                joining_date,
                plan_info,
                total_due,
                pt_info,
                full_address
            ]

            for col_num, cell_value in enumerate(row_data, 1):
                cell = worksheet.cell(row=index + 1, column=col_num)
                if 'http' in str(cell_value):
                    cell.hyperlink = cell_value
                    cell.value = "View Photo"
                    cell.style = "Hyperlink"
                else:
                    cell.value = cell_value
                cell.alignment = Alignment(vertical='center', wrap_text=True)

        workbook.save(response)
        return response

    # Paginate
    paginator = Paginator(all_members_data, 10)
    page_number = request.GET.get('page')
    members = paginator.get_page(page_number)

    # Final pass for display data only
    for member in members:
        histories = member.all_histories_list
        membership_due_amount = sum((mh.total_amount - mh.paid_amount) for mh in histories if mh.status == 'active')
        pt_due_amount = sum((pt.total_amount - pt.paid_amount) for pt in member.personal_trainer.all() if pt.status == 'active')
        member.total_due = membership_due_amount + pt_due_amount

        first_history = min(histories, key=lambda mh: mh.membership_start_date, default=None)
        member.join_date = first_history.membership_start_date if first_history else None

    return render(request, 'members/member_list.html', {'members': members})

@require_POST
@login_required(login_url='login')
@custom_permission_required('change_member')
def toggle_member_status(request, member_id):
    gym = getattr(request, 'gym', None)
    member = Member.objects.filter(id=member_id, gym=gym).first()
    if not member:
        return JsonResponse({'status': 'error', 'message': 'Member not found.'}, status=404)
    if member.status == 'active':
        member.status = 'inactive'
    else:
        member.status = 'active'
    member.save()
    messages.success(request, f'Member status has been updated to {member.status}.')
    return JsonResponse({'status': 'success', 'message': 'Member status updated successfully.', 'new_status': member.status})

@require_POST
@login_required(login_url='login')
@custom_permission_required('delete_member')
def delete_member(request, member_id):
    gym = getattr(request, 'gym', None)
    member = Member.objects.filter(id=member_id, gym=gym).first()
    if not member:
        return JsonResponse({'status': 'error', 'message': 'Member not found.'}, status=404)
    try:
        with transaction.atomic():
            member.delete()
        messages.success(request, 'Member has been deleted successfully.')
        return JsonResponse({'status': 'success', 'message': 'Member has been deleted successfully.'})
    except IntegrityError:
        # If deletion is protected, mark as deleted
        member.is_deleted = True
        member.save()
        messages.success(request, 'Member has been marked as deleted due to existing dependencies.')
        return JsonResponse({'status': 'success', 'message': 'Member has been marked as deleted due to existing dependencies.'})
    except Exception as e:
        messages.error(request, f'An error occurred: {e}')
        return JsonResponse({'status': 'error', 'message': str(e)}, status=500)

@never_cache
@login_required(login_url='login')
@custom_permission_required('change_member')
def assign_membership_plan(request, member_id, history_id=None):
    gym = getattr(request, 'gym', None)
    member = Member.objects.filter(id=member_id, gym=gym).first()
    if not member:
        messages.error(request, 'Member not found.')
        return redirect('member_list')

    # Fetch existing history if history_id is provided
    history_instance = None
    if history_id:
        history_instance = MembershipHistory.objects.filter(id=history_id, member=member, gym=gym).first()
        if not history_instance:
            messages.error(request, 'Membership record not found.')
            return redirect('member_profile', member_id=member.id)
   
    plans = MembershipPlan.objects.filter(gym=gym)
    plans_json = serialize('json', plans)
    gym_json = serialize('json', [gym])

    if request.method == 'POST':
        form = MembershipHistoryForm(request.POST, instance=history_instance, gym=gym)
        if form.is_valid():
            history = form.save(commit=False)
            history.member = member
            history.gym = gym
            history.transaction_id = request.POST.get('transaction_id')

            # Calculate total amount with GST
            plan_fee = history.plan.offer_price
            registration_fee = form.cleaned_data.get('registration_fee', 0) or 0
            discount = form.cleaned_data.get('discount', 0) or 0
            
            # Calculate the subtotal before GST
            subtotal = plan_fee + registration_fee - discount
            
            # Calculate GST if enabled
            gst_amount = Decimal(0)
            sgst = Decimal(0)
            cgst = Decimal(0)
            
            if gym.gst_enabled:
                gst_rate = gym.gst_rate
                gst_amount = (subtotal * gst_rate) / 100
                sgst = gst_amount / 2
                cgst = gst_amount / 2

            # Calculate the final total amount
            total_amount = subtotal + gst_amount
            
            # Save the values to the history instance
            history.sgst = sgst.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
            history.cgst = cgst.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
            history.gst_amount = gst_amount.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
            history.total_amount = total_amount.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

            history.save()

            if form.cleaned_data.get('follow_up_date'):
                follow_up_date = form.cleaned_data['follow_up_date']
                MembershipHistory.objects.filter(member=member, status='active', gym=gym).exclude(paid_amount=F('total_amount')).update(follow_up_date=follow_up_date)
                PersonalTrainer.objects.filter(member=member, status='active', gym=gym).exclude(paid_amount=F('total_amount')).update(follow_up_date=follow_up_date)

            # Handle payment record
            if history.paid_amount > 0:
                payment, created = Payment.objects.update_or_create(
                    membership_history=history,
                    defaults={
                        'gym': gym,
                        'member': member,
                        'amount': history.paid_amount,
                        'payment_mode': history.payment_mode,
                        'transaction_id': history.transaction_id,
                        'comment': f"Payment for {history.plan.title} (Updated)" if history_id else f"Initial payment for {history.plan.title}",
                        'payment_date': history.payment_date
                    }
                )

            member.membership_plan = history.plan
            member.save()
            
            # Dispatch automated in-app platform notifications & popups
            try:
                notify_membership_activated(member=member, history=history, is_upgrade=bool(history_id))
                if history.paid_amount > 0 and 'payment' in locals() and payment:
                    notify_payment_submitted(payment=payment, member=member, invoice=history, invoice_type='membership')
            except Exception as notif_err:
                pass

            action_text = "upgraded" if history_id else "assigned"
            # Send WhatsApp notification on plan assignment / renewal
            if gym.whatsapp_enabled:
                try:
                    whatsapp_service = WhatsAppService(gym_id=gym.id)
                    wa_res = whatsapp_service.send_membership_plan_details(
                        request=request,
                        to_number=member.mobile_number,
                        name=member.name,
                        gym_name=gym.name,
                        plan_name=history.plan.title,
                        amount_paid=history.paid_amount,
                        balance=history.total_amount - history.paid_amount,
                        expiry_date=history.get_end_date().strftime('%d/%m/%Y'),
                        gym_contact_number=gym.phone,
                        gym_logo_url=gym.logo.url if gym.logo else None
                    )
                    msg_text = f'Membership plan "{history.plan.title}" {action_text} successfully.'
                    if wa_res and wa_res.get('success'):
                        messages.success(request, msg_text, extra_tags='whatsapp_success')
                    else:
                        messages.warning(request, msg_text, extra_tags='whatsapp_warning')
                except Exception:
                    messages.warning(request, msg_text, extra_tags='whatsapp_warning')
            else:
                msg_text = f'Membership plan "{history.plan.title}" {action_text} successfully.'
                messages.warning(request, msg_text, extra_tags='whatsapp_warning')

            return redirect('billing:invoice', member_id=member.id, history_id=history.id)
    else:
        initial_data = {}
        if not history_id:
            # For new assignments, pre-fill start date based on previous plan's expiry
            latest_history = MembershipHistory.objects.filter(member=member, gym=gym, is_deleted=False).order_by('-membership_start_date').first()
            initial_start_date = timezone.localdate()
            if latest_history:
                end_date = latest_history.get_end_date()
                if end_date:
                    # Start the new plan the day after the old one expires
                    initial_start_date = end_date + timedelta(days=1)
            initial_data['membership_start_date'] = initial_start_date

        form = MembershipHistoryForm(instance=history_instance, gym=gym, initial=initial_data)
    return render(request, 'members/membership_plan_assign.html', {
        'member': member, 
        'plans': plans, 
        'plans_json': plans_json, 
        'form': form, 
        'gym_json': gym_json,
        'history': history_instance
    })

@never_cache
@login_required(login_url='login')
@custom_permission_required('change_member')
def assign_pt_trainer(request, member_id):
    gym = getattr(request, 'gym', None)
    member = Member.objects.filter(id=member_id, gym=gym).first()
    if not member:
        messages.error(request, 'Member not found.')
        return redirect('member_list')
    
    # Filter trainers by the current gym
    trainers = Trainer.objects.filter(gym=gym)
    trainers_json = serialize('json', trainers)
    
    if request.method == 'POST':
        form = PersonalTrainerForm(request.POST, gym=gym) # Pass gym to the form
        if form.is_valid():
            pt_assignment = form.save(commit=False)
            pt_assignment.member = member
            pt_assignment.gym = gym # Assign gym to the instance
            pt_assignment.transaction_id = request.POST.get('transaction_id')
            pt_assignment.save()

            payment = None
            if pt_assignment.paid_amount > 0:
                payment = Payment.objects.create(
                    gym=gym,
                    member=member,
                    amount=pt_assignment.paid_amount,
                    payment_mode=pt_assignment.payment_mode,
                    transaction_id=pt_assignment.transaction_id,
                    comment=f"Initial payment for Personal Trainer: {pt_assignment.trainer.name}",
                    personal_trainer=pt_assignment,
                    payment_date=pt_assignment.payment_date
                )

            # Dispatch automated in-app platform notifications
            try:
                notify_pt_assigned(member=member, pt_assignment=pt_assignment)
                if payment:
                    notify_payment_submitted(payment=payment, member=member, invoice=pt_assignment, invoice_type='pt')
            except Exception as notif_err:
                pass

            messages.success(request, f'Personal Trainer "{pt_assignment.trainer.name}" assigned to {member.first_name} {member.last_name}.')
            return redirect('billing:pt_invoice', member_id=member.id, pt_invoice_id=pt_assignment.id)
    else:
        form = PersonalTrainerForm(gym=gym) # Pass gym to the form
    
    return render(request, 'members/assign_PT_trainer.html', {'member': member, 'trainers': trainers, 'trainers_json': trainers_json, 'form': form})


@login_required(login_url='login')
@custom_permission_required('change_member')
def unfreeze_membership(request, membership_id):
    gym = getattr(request, 'gym', None)
    membership = MembershipHistory.objects.filter(id=membership_id, gym=gym).first()
    if not membership:
        messages.error(request, 'Membership record not found.')
        return redirect('member_list')
    
    if request.method == 'POST':
        freeze = MembershipFreeze.objects.filter(membership=membership, unfreeze_date__isnull=True).first()
        if freeze:
            freeze.unfreeze_date = timezone.localdate()
            freeze.save()
            
            membership.status = 'active'
            membership.save()
            
            messages.success(request, "Membership has been unfrozen.")
        else:
            messages.error(request, "No active freeze found for this membership.")
        
        return redirect('member_profile', member_id=membership.member.id)
    
    return redirect('member_profile', member_id=membership.member.id)


@login_required(login_url='login')
@custom_permission_required('change_member')
def freeze_membership(request, membership_id):
    gym = getattr(request, 'gym', None)
    membership = MembershipHistory.objects.filter(id=membership_id, gym=gym).first()
    if not membership:
        messages.error(request, 'Membership record not found.')
        return redirect('member_list')
    
    if request.method == 'POST':
        reason = request.POST.get('reason', '')
        
        # Check if the membership is already frozen
        if membership.is_frozen:
            messages.error(request, "This membership is already frozen.")
            return redirect('member_profile', member_id=membership.member.id)

        MembershipFreeze.objects.create(
            membership=membership,
            freeze_date=timezone.localdate(),
            reason=reason
        )
        
        membership.status = 'frozen'
        membership.save()
        
        messages.success(request, "Membership has been frozen.")
        return redirect('member_profile', member_id=membership.member.id)
    
    return redirect('member_profile', member_id=membership.member.id)


@login_required(login_url='login')
@custom_permission_required(['change_member', 'change_dietplan', 'add_dietplan'])
def assign_diet_plan(request, member_id):
    gym = getattr(request, 'gym', None)
    member = Member.objects.filter(id=member_id, gym=gym).first()
    if not member:
        messages.error(request, 'Member not found.')
        return redirect('member_list')
    
    if request.method == 'POST':
        form = AssignDietPlanForm(request.POST, gym=gym)
        if form.is_valid():
            assign_diet_plan = form.save(commit=False)
            assign_diet_plan.member = member
            assign_diet_plan.gym = gym
            assign_diet_plan.save()
            
            # Generate public link and send WhatsApp message
            if gym.whatsapp_enabled:
                try:
                    public_link = request.build_absolute_uri(f'/management/public/diet-plan/{assign_diet_plan.diet_plan.pk}/')
                    
                    whatsapp_msg = (
                        f"Hello {member.first_name},\n\n"
                        f"A new Diet Plan *{assign_diet_plan.diet_plan.name}* has been assigned to you by *{gym.name}*.\n\n"
                        f"You can view your detailed plan and download any attached documents using the secure link below:\n\n"
                        f"🔗 {public_link}\n\n"
                        f"Stay healthy and keep progressing!\n"
                        f"- Team {gym.name}"
                    )
                    
                    whatsapp_service = WhatsAppService(gym_id=gym.id)
                    wa_res = whatsapp_service.send_message(to_number=member.mobile_number, message=whatsapp_msg)
                    msg_text = f'Diet plan assigned to {member.name} successfully.'
                    if wa_res and wa_res.get('success'):
                        messages.success(request, msg_text, extra_tags='whatsapp_success')
                    else:
                        messages.warning(request, msg_text, extra_tags='whatsapp_warning')
                except Exception:
                    msg_text = f'Diet plan assigned to {member.name} successfully.'
                    messages.warning(request, msg_text, extra_tags='whatsapp_warning')
            else:
                msg_text = f'Diet plan assigned to {member.name} successfully.'
                messages.warning(request, msg_text, extra_tags='whatsapp_warning')
            
            return redirect('member_profile', member_id=member.id)
    else:
        form = AssignDietPlanForm(gym=gym)
    
    return render(request, 'members/assign_diet_plan.html', {'member': member, 'form': form})


@login_required(login_url='login')
@custom_permission_required(['change_member', 'change_workoutplan', 'add_workoutplan'])
def assign_workout_plan(request, member_id):
    gym = getattr(request, 'gym', None)
    member = Member.objects.filter(id=member_id, gym=gym).first()
    if not member:
        messages.error(request, 'Member not found.')
        return redirect('member_list')
    
    if request.method == 'POST':
        form = AssignWorkoutPlanForm(request.POST, gym=gym)
        if form.is_valid():
            assign_workout_plan = form.save(commit=False)
            assign_workout_plan.member = member
            assign_workout_plan.gym = gym
            assign_workout_plan.save()
            
            # Generate public link and send WhatsApp message
            if gym.whatsapp_enabled:
                try:
                    public_link = request.build_absolute_uri(f'/management/public/workout-plan/{assign_workout_plan.workout_plan.pk}/')
                    
                    whatsapp_msg = (
                        f"Hello {member.first_name},\n\n"
                        f"A new Workout Plan *{assign_workout_plan.workout_plan.name}* has been assigned to you by *{gym.name}*.\n\n"
                        f"You can view your detailed workout routine using the secure link below:\n\n"
                        f"🔗 {public_link}\n\n"
                        f"Stay strong and keep pushing!\n"
                        f"- Team {gym.name}"
                    )
                    
                    whatsapp_service = WhatsAppService(gym_id=gym.id)
                    wa_res = whatsapp_service.send_message(to_number=member.mobile_number, message=whatsapp_msg)
                    msg_text = f'Workout plan assigned to {member.name} successfully.'
                    if wa_res and wa_res.get('success'):
                        messages.success(request, msg_text, extra_tags='whatsapp_success')
                    else:
                        messages.warning(request, msg_text, extra_tags='whatsapp_warning')
                except Exception:
                    msg_text = f'Workout plan assigned to {member.name} successfully.'
                    messages.warning(request, msg_text, extra_tags='whatsapp_warning')
            else:
                msg_text = f'Workout plan assigned to {member.name} successfully.'
                messages.warning(request, msg_text, extra_tags='whatsapp_warning')
            
            return redirect('member_profile', member_id=member.id)
    else:
        form = AssignWorkoutPlanForm(gym=gym)
    
    return render(request, 'members/assign_workout_plan.html', {'member': member, 'form': form})

@never_cache
@login_required(login_url='login')
@custom_permission_required('change_member')
@require_POST
def delete_assigned_diet_plan(request, assigned_plan_id):
    gym = getattr(request, 'gym', None)
    assigned_plan = AssignDietPlan.objects.filter(id=assigned_plan_id, gym=gym).first()
    if not assigned_plan:
        return JsonResponse({'status': 'error', 'message': 'Diet plan assignment not found.'}, status=404)
    try:
        assigned_plan.delete()
        return JsonResponse({'status': 'success', 'message': 'Assigned diet plan has been deleted.'})
    except Exception as e:
        return JsonResponse({'status': 'error', 'message': str(e)}, status=500)

@never_cache
@login_required(login_url='login')
@custom_permission_required('change_member')
@require_POST
def delete_assigned_workout_plan(request, assigned_plan_id):
    gym = getattr(request, 'gym', None)
    assigned_plan = AssignWorkoutPlan.objects.filter(id=assigned_plan_id, gym=gym).first()
    if not assigned_plan:
        return JsonResponse({'status': 'error', 'message': 'Workout plan assignment not found.'}, status=404)
    try:
        assigned_plan.delete()
        return JsonResponse({'status': 'success', 'message': 'Assigned workout plan has been deleted.'})
    except Exception as e:
        return JsonResponse({'status': 'error', 'message': str(e)}, status=500)


@never_cache
@login_required(login_url='login')
@custom_permission_required('change_member')
@require_POST
def reset_member_password(request, member_id):
    gym = getattr(request, 'gym', None)
    member = Member.objects.filter(id=member_id, gym=gym, is_deleted=False).first()
    if not member:
        return JsonResponse({'status': 'error', 'message': 'Member not found.'}, status=404)
    
    custom_pwd = request.POST.get('new_password', '').strip()
    if not custom_pwd:
        import random
        phone_suffix = member.mobile_number[-4:] if member.mobile_number and len(member.mobile_number) >= 4 else str(random.randint(1000, 9999))
        custom_pwd = f"Fit@{phone_suffix}"
        
    user, raw_pwd = member.create_or_update_user_account(raw_password=custom_pwd)
    
    login_url = request.build_absolute_uri('/login/')
    share_text = f"*ONEX BOOST Gym Password Reset*\n\nHi {member.name},\nYour account password has been reset successfully.\n\n• *Username / ID:* {member.member_id}\n• *Mobile:* {member.mobile_number}\n• *New Password:* {raw_pwd}\n• *Login URL:* {login_url}\n\nYou can log in using your Mobile Number or Username."
    
    return JsonResponse({
        'status': 'success',
        'message': f'Password for {member.name} has been reset successfully!',
        'name': member.name,
        'member_id': member.member_id,
        'username': user.username,
        'mobile': member.mobile_number,
        'new_password': raw_pwd,
        'share_text': share_text,
    })
